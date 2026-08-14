"""热梗库导入与最终采用记录。"""

import io
import unittest
from types import SimpleNamespace
from unittest.mock import patch

from app.api.meme_library import can_manage_meme_entry
from app.services.llm_client import LLMConfig
from app.services.meme_library import (
    MEME_LIBRARY_HEADERS,
    build_meme_library_workbook,
    build_meme_retrieval_text,
    meme_embedding_model_key,
    normalize_manual_meme_entry,
    parse_meme_library_file,
)
from app.services.meme_rag import (
    MAX_MEME_RERANK_CANDIDATES,
    MIN_MEME_SCENE_FIT_SCORE,
    MIN_MEME_VECTOR_SIMILARITY,
    apply_meme_scene_fit_results,
    build_chapter_meme_query,
    build_final_meme_usage,
    collect_adopted_meme_phrases,
    find_repeated_meme_phrases,
    find_reused_event_meme_phrases,
    meets_meme_relevance_threshold,
)
from app.services.meme_usage_enforcer import (
    ensure_minimum_meme_usage,
    repair_repeated_meme_usage,
)


class MemeLibraryTests(unittest.TestCase):
    def test_export_workbook_contains_all_import_fields(self) -> None:
        from openpyxl import load_workbook

        entries = [
            SimpleNamespace(
                phrase="邪修",
                meaning="不按常规但意外高效的做法",
                suitable_scenes="熟人轻松调侃",
            ),
            SimpleNamespace(
                phrase="预制××",
                meaning="批量套模板的事物",
                suitable_scenes="吐槽缺少个性",
            ),
        ]

        content = build_meme_library_workbook(entries)
        workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
        rows = list(workbook.active.iter_rows(values_only=True))

        self.assertEqual(rows[0], MEME_LIBRARY_HEADERS)
        self.assertEqual(len(rows), 3)
        self.assertEqual(rows[1][0], "邪修")
        self.assertEqual(rows[1][2], "熟人轻松调侃")

    def test_builtin_entries_are_manageable_but_user_entries_remain_private(self) -> None:
        current_user_id = "00000000-0000-0000-0000-000000000001"
        builtin = SimpleNamespace(namespace="builtin", owner_id=None)
        own_entry = SimpleNamespace(namespace=f"user:{current_user_id}", owner_id=current_user_id)
        other_entry = SimpleNamespace(
            namespace="user:00000000-0000-0000-0000-000000000002",
            owner_id="00000000-0000-0000-0000-000000000002",
        )

        self.assertTrue(can_manage_meme_entry(builtin, current_user_id))
        self.assertTrue(can_manage_meme_entry(own_entry, current_user_id))
        self.assertFalse(can_manage_meme_entry(other_entry, current_user_id))

    def test_embedding_text_only_uses_three_semantic_fit_fields(self) -> None:
        retrieval_text = build_meme_retrieval_text(
            {
                "phrase": "邪修",
                "meaning": "不按常规但意外高效的做法",
                "suitable_scenes": "熟人轻松调侃一种省事办法",
            }
        )

        self.assertIn("邪修", retrieval_text)
        self.assertIn("意外高效", retrieval_text)
        self.assertIn("熟人轻松调侃", retrieval_text)
        self.assertTrue(meme_embedding_model_key("qwen-test").endswith("::meme-fit-v2"))

    def test_chapter_query_uses_only_active_character_profiles(self) -> None:
        novel = SimpleNamespace(genre="都市")
        context = {
            "target": {
                "task_input": {
                    "chapter_plan": {
                        "participants": ["林安", "欧阳果"],
                        "core_event": "两人因临时误会互相试探",
                        "dramatic_turn": "欧阳果先认错，反而让林安更愧疚",
                        "reader_payoff": "关系推进",
                    }
                }
            },
            "constraints": {
                "characters": [
                    {"name": "林安", "occupation": "产品经理"},
                    {"name": "欧阳果", "occupation": "插画师"},
                    {"name": "未出场角色", "occupation": "律师"},
                ]
            },
            "recent_chapters": [],
        }

        query = build_chapter_meme_query(novel, context)

        self.assertIn("林安", query)
        self.assertIn("欧阳果", query)
        self.assertIn("关系推进", query)
        self.assertNotIn("未出场角色", query)
        self.assertNotIn("律师", query)

    def test_scene_fit_requires_meaning_relationship_emotion_and_full_micro_scene(self) -> None:
        candidates = [{"entry_id": "one", "phrase": "邪修", "retrieval_score": 0.08}]
        base_result = {
            "entry_id": "one",
            "eligible": True,
            "fit_score": MIN_MEME_SCENE_FIT_SCORE,
            "meaning_fit": True,
            "relationship_fit": False,
            "emotion_fit": True,
            "speaker": "林安",
            "listener": "欧阳果",
            "relationship": "熟悉的朋友",
            "emotion": "轻松调侃",
            "speech_act": "评价对方的省事办法",
            "required_setup": "欧阳果展示了一种非常规但高效的做法",
            "scene_anchor": "林安看完结果",
            "response": "欧阳果顺势承认并反问",
            "plot_consequence": "两人决定继续采用该办法",
            "reason": "含义一致",
        }

        rejected, rejected_count = apply_meme_scene_fit_results(
            candidates,
            {"results": [base_result]},
        )
        accepted, accepted_rejected_count = apply_meme_scene_fit_results(
            candidates,
            {"results": [{**base_result, "relationship_fit": True}]},
        )

        self.assertEqual(rejected, [])
        self.assertEqual(rejected_count, 1)
        self.assertEqual(accepted_rejected_count, 0)
        self.assertEqual(accepted[0]["scene_fit"]["speech_act"], "评价对方的省事办法")

    def test_event_meme_can_only_be_adopted_once(self) -> None:
        chapters = [
            SimpleNamespace(
                context_snapshot={
                    "meme_reference": {
                        "usage": {
                            "adopted": [
                                {"phrase": "邪修"},
                                {"phrase": "预制××"},
                            ]
                        }
                    }
                }
            ),
            SimpleNamespace(
                context_snapshot={
                    "meme_reference": {
                        "usage": {"adopted": [{"phrase": "邪修"}]}
                    }
                }
            ),
        ]

        used = collect_adopted_meme_phrases(chapters)
        reused = find_reused_event_meme_phrases(
            "他笑道：“你这是预制友情吧。”",
            used,
        )

        self.assertEqual(used, ["邪修", "预制××"])
        self.assertEqual(reused, ["预制××"])
        self.assertEqual(
            find_repeated_meme_phrases(
                "这做法很邪修。确实邪修。",
                used,
            ),
            ["邪修"],
        )

    def test_meme_recall_uses_conservative_vector_floor(self) -> None:
        self.assertTrue(
            meets_meme_relevance_threshold(
                vector_similarity=MIN_MEME_VECTOR_SIMILARITY,
            )
        )
        self.assertFalse(
            meets_meme_relevance_threshold(
                vector_similarity=MIN_MEME_VECTOR_SIMILARITY - 0.001,
            )
        )
        self.assertEqual(MIN_MEME_VECTOR_SIMILARITY, 0.40)
        self.assertEqual(MAX_MEME_RERANK_CANDIDATES, 12)

    def test_imports_three_column_utf8_csv(self) -> None:
        csv_text = (
            ",".join(MEME_LIBRARY_HEADERS)
            + "\n"
            + "示例表达,准确含义,朋友之间轻松调侃\n"
        )
        entries, errors = parse_meme_library_file("memes.csv", csv_text.encode("utf-8"))

        self.assertEqual(errors, [])
        self.assertEqual(entries[0]["phrase"], "示例表达")
        self.assertIn("适用人物关系", entries[0]["retrieval_text"])

    def test_rejects_incomplete_rows(self) -> None:
        csv_text = (
            ",".join(MEME_LIBRARY_HEADERS)
            + "\n"
            + "示例表达,准确含义,\n"
        )
        entries, errors = parse_meme_library_file("memes.csv", csv_text.encode("utf-8"))

        self.assertEqual(entries, [])
        self.assertIn("适用场景", errors[0])

    def test_manual_entry_uses_same_validation_and_index_fields_as_import(self) -> None:
        entry = normalize_manual_meme_entry(
            {
                "phrase": "邪修",
                "meaning": "不按常规但意外高效的做法",
                "suitable_scenes": "熟人轻松调侃对方的省事办法",
            }
        )

        self.assertEqual(entry["normalized_phrase"], "邪修")
        self.assertIn("熟人轻松调侃", entry["retrieval_text"])
        self.assertEqual(len(entry["content_hash"]), 64)

    def test_final_usage_only_keeps_phrases_present_after_review(self) -> None:
        pack = {
            "references": [
                {"entry_id": "one", "phrase": "低山臭水遇知音", "source_type": "builtin"},
                {"entry_id": "two", "phrase": "敬自己一杯", "source_type": "builtin"},
            ]
        }
        progress = {
            "meme_usage_plan": [
                {
                    "phrase": "低山臭水遇知音",
                    "decision": "use",
                    "speaker": "林安",
                    "scene_anchor": "朋友发现共同怪笑点",
                }
            ]
        }

        result = build_final_meme_usage(pack, progress, "林安笑了：“这算低山臭水遇知音。”")

        self.assertEqual(result["adopted_count"], 1)
        self.assertEqual(result["adopted"][0]["speaker"], "林安")
        self.assertEqual(result["adopted"][0]["phrase"], "低山臭水遇知音")

    def test_final_usage_does_not_count_unplanned_text_match(self) -> None:
        pack = {
            "references": [
                {"entry_id": "one", "phrase": "预制××", "source_type": "builtin"},
            ]
        }

        result = build_final_meme_usage(
            pack,
            {"meme_usage_plan": []},
            "中介发来的是一套批量复制的预制房源信息。",
        )

        self.assertEqual(result["adopted_count"], 0)
        self.assertEqual(result["adopted"], [])

    def test_enforcer_does_not_call_llm_when_meme_already_present(self) -> None:
        pack = {
            "references": [
                {"entry_id": "one", "phrase": "邪修", "source_type": "builtin"},
            ]
        }
        content = "林安看完那套省事做法，笑道：“你这算生活邪修。”"

        with patch(
            "app.services.meme_usage_enforcer.LLMClient.complete_json"
        ) as mocked_complete:
            result = ensure_minimum_meme_usage(
                content=content,
                reference_pack=pack,
                chapter_progress={
                    "meme_usage_plan": [
                        {"phrase": "邪修", "decision": "use"}
                    ]
                },
                llm_config=LLMConfig(
                    base_url="https://example.invalid/v1",
                    api_key="test",
                    model="test",
                ),
            )

        self.assertEqual(result["status"], "already_satisfied")
        self.assertEqual(result["usage"]["adopted_count"], 1)
        mocked_complete.assert_not_called()

    def test_enforcer_never_injects_a_meme_after_drafting(self) -> None:
        pack = {
            "references": [
                {
                    "entry_id": "one",
                    "phrase": "邪修",
                    "meaning": "不按常规但高效的办法",
                    "suitable_scenes": "熟人调侃生活技巧",
                    "source_type": "builtin",
                },
            ]
        }
        with patch(
            "app.services.meme_usage_enforcer.LLMClient.complete_json"
        ) as mocked_complete:
            result = ensure_minimum_meme_usage(
                content="林安看着她把鸡蛋直接磕进饭盒，忍不住笑了。\n\n她认真记下火候。",
                reference_pack=pack,
                chapter_progress={
                    "meme_usage_plan": [
                        {"phrase": "邪修", "decision": "skip", "reason": "初稿未采用"}
                    ]
                },
                llm_config=LLMConfig(
                    base_url="https://example.invalid/v1",
                    api_key="test",
                    model="test",
                ),
            )

        self.assertEqual(result["status"], "optional_skipped")
        self.assertEqual(result["usage"]["adopted_count"], 0)
        self.assertNotIn("邪修", result["content"])
        mocked_complete.assert_not_called()

    def test_enforcer_does_not_force_an_inapplicable_candidate(self) -> None:
        pack = {
            "references": [
                {
                    "entry_id": "one",
                    "phrase": "邪修",
                    "meaning": "不按常规但高效的办法",
                    "suitable_scenes": "熟人轻松调侃",
                    "scene_fit": {
                        "speaker": "林安",
                        "speech_act": "调侃",
                        "required_setup": "先看到非常规办法",
                        "scene_anchor": "当面交流",
                        "response": "对方接话",
                        "plot_consequence": "共同改变做法",
                    },
                }
            ]
        }
        with patch(
            "app.services.meme_usage_enforcer.LLMClient.complete_json"
        ) as mocked_complete:
            result = ensure_minimum_meme_usage(
                content="他独自在空房间里查完资料，关灯离开。",
                reference_pack=pack,
                chapter_progress={},
                llm_config=LLMConfig(
                    base_url="https://example.invalid/v1",
                    api_key="test",
                    model="test",
                ),
            )

        self.assertEqual(result["status"], "optional_skipped")
        self.assertEqual(result["content"], "他独自在空房间里查完资料，关灯离开。")
        mocked_complete.assert_not_called()

    def test_repeated_template_meme_is_repaired_with_one_local_patch(self) -> None:
        content = (
            "林予安捏着规则纸：“这份东西有点预制生活的感觉。”\n\n"
            "沈栀夏抿住嘴：“预制不是贬义，至少提前考虑了风险。”\n\n"
            "两人继续删减规则。"
        )
        response = {
            "patches": [
                {
                    "paragraph_index": 2,
                    "new_text": "沈栀夏抿住嘴：“提前安排也不是坏事，至少考虑了风险。”",
                    "reason": "保留她的反驳意图，去掉第二次模板表达",
                }
            ]
        }
        with patch(
            "app.services.meme_usage_enforcer.LLMClient.complete_json",
            return_value=("{}", response),
        ):
            result = repair_repeated_meme_usage(
                content=content,
                phrases=["预制××"],
                llm_config=LLMConfig(
                    base_url="https://example.invalid/v1",
                    api_key="test",
                    model="test",
                ),
            )

        self.assertEqual(result["status"], "repaired")
        self.assertEqual(result["paragraph_indexes"], [2])
        self.assertIn("预制生活", result["content"])
        self.assertNotIn("预制不是", result["content"])

    def test_repeated_meme_repair_retries_with_validation_error(self) -> None:
        content = "“今晚资源有限，能当个事儿办的都得当个事儿办。”"
        responses = [
            (
                "{}",
                {
                    "patches": [
                        {
                            "paragraph_index": 1,
                            "new_text": content.strip("“”"),
                            "reason": "第一次没有正确去重",
                        }
                    ]
                },
            ),
            (
                "{}",
                {
                    "patches": [
                        {
                            "paragraph_index": 1,
                            "new_text": "“今晚资源有限，这事我会当个事儿办。”",
                            "reason": "保留一次热梗并去掉回声重复",
                        }
                    ]
                },
            ),
        ]
        with patch(
            "app.services.meme_usage_enforcer.LLMClient.complete_json",
            side_effect=responses,
        ) as complete_json:
            result = repair_repeated_meme_usage(
                content=content,
                phrases=["当个事儿办"],
                llm_config=LLMConfig(
                    base_url="https://example.invalid/v1",
                    api_key="test",
                    model="test",
                ),
            )

        self.assertEqual(result["status"], "repaired")
        self.assertEqual(result["attempt"], 2)
        self.assertEqual(result["content"].count("当个事儿办"), 1)
        self.assertEqual(complete_json.call_count, 2)

    def test_repeated_meme_repair_keeps_original_when_patch_is_incomplete(self) -> None:
        content = "这做法很邪修。\n\n他又说了一遍：“确实邪修。”"
        with patch(
            "app.services.meme_usage_enforcer.LLMClient.complete_json",
            return_value=("{}", {"patches": []}),
        ):
            result = repair_repeated_meme_usage(
                content=content,
                phrases=["邪修"],
                llm_config=LLMConfig(
                    base_url="https://example.invalid/v1",
                    api_key="test",
                    model="test",
                ),
            )

        self.assertEqual(result["status"], "failed")
        self.assertEqual(result["content"], content)


if __name__ == "__main__":
    unittest.main()
