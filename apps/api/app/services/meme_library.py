"""热梗库导入、版本同步和 Qwen 向量索引。"""

from __future__ import annotations

import csv
import hashlib
import io
import re
from pathlib import Path
from typing import Any
from uuid import UUID

from openpyxl import load_workbook
from sqlalchemy import or_, select
from sqlalchemy.orm import Session

from app.core.config import settings
from app.models.meme_entry import MemeEntry
from app.models.user import User
from app.services.embedding_client import EmbeddingClient, build_embedding_config


MEME_LIBRARY_HEADERS = ("热梗", "含义", "出处事件", "适用场景", "流行时间", "来源链接")
MAX_MEME_IMPORT_BYTES = 8 * 1024 * 1024
MAX_MEME_IMPORT_ROWS = 2000
MEME_EMBEDDING_SCHEMA = "meme-fit-v2"


def meme_embedding_model_key(model: str) -> str:
    """Version meme vectors separately so retrieval-text changes trigger reindexing."""
    return f"{str(model or '').strip()}::{MEME_EMBEDDING_SCHEMA}"


def normalize_meme_phrase(value: Any) -> str:
    return re.sub(r"\s+", "", str(value or "").strip()).casefold()[:120]


def parse_popularity_years(value: Any) -> tuple[int | None, int | None]:
    years = [int(item) for item in re.findall(r"(?<!\d)(20\d{2})(?!\d)", str(value or ""))]
    if not years:
        return None, None
    return min(years), max(years)


def build_meme_retrieval_text(entry: dict[str, Any]) -> str:
    """Only embed the three fields that determine whether a meme can be used."""
    return "\n".join(
        [
            f"网络表达：{entry['phrase']}",
            f"真实含义：{entry['meaning']}",
            f"适用人物关系、情绪与场景：{entry['suitable_scenes']}",
        ]
    )[:1400]


def build_meme_content_hash(entry: dict[str, Any]) -> str:
    canonical = "\n".join(
        str(entry.get(key) or "").strip()
        for key in (
            "phrase",
            "meaning",
            "origin_event",
            "suitable_scenes",
            "popularity_period",
            "source_urls",
        )
    )
    return hashlib.sha256(canonical.encode("utf-8")).hexdigest()


def _clean_source_urls(value: Any) -> list[str]:
    urls = []
    raw_values = value if isinstance(value, (list, tuple)) else str(value or "").split("|")
    for raw in raw_values:
        url = raw.strip()
        if not url:
            continue
        if not re.match(r"^https?://", url, flags=re.IGNORECASE):
            raise ValueError(f"来源链接必须以 http:// 或 https:// 开头：{url}")
        if url not in urls:
            urls.append(url)
    if not urls:
        raise ValueError("来源链接不能为空")
    return urls[:5]


def normalize_manual_meme_entry(payload: dict[str, Any]) -> dict[str, Any]:
    """复用导入规则校验单条手工录入，并生成检索文本与内容版本。"""
    entry = _normalize_import_row(
        {
            "热梗": payload.get("phrase"),
            "含义": payload.get("meaning"),
            "出处事件": payload.get("origin_event"),
            "适用场景": payload.get("suitable_scenes"),
            "流行时间": payload.get("popularity_period"),
            "来源链接": payload.get("source_urls"),
        },
        1,
    )
    entry["retrieval_text"] = build_meme_retrieval_text(entry)
    entry["content_hash"] = build_meme_content_hash(entry)
    return entry


def _normalize_import_row(row: dict[str, Any], row_number: int) -> dict[str, Any]:
    source_urls_value = row.get("来源链接")
    values = {
        header: (
            "|".join(str(item).strip() for item in source_urls_value)
            if header == "来源链接" and isinstance(source_urls_value, (list, tuple))
            else str(row.get(header) or "").strip()
        )
        for header in MEME_LIBRARY_HEADERS
    }
    missing = [header for header, value in values.items() if not value]
    if missing:
        raise ValueError(f"第 {row_number} 行缺少：{'、'.join(missing)}")
    if len(values["热梗"]) > 120:
        raise ValueError(f"第 {row_number} 行热梗超过 120 字")
    start_year, end_year = parse_popularity_years(values["流行时间"])
    return {
        "phrase": values["热梗"],
        "normalized_phrase": normalize_meme_phrase(values["热梗"]),
        "meaning": values["含义"][:1000],
        "origin_event": values["出处事件"][:1600],
        "suitable_scenes": values["适用场景"][:1200],
        "popularity_period": values["流行时间"][:80],
        "popularity_year_start": start_year,
        "popularity_year_end": end_year,
        "source_urls": _clean_source_urls(source_urls_value),
    }


def _xlsx_rows(content: bytes) -> list[dict[str, Any]]:
    workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    sheet = workbook.active
    rows = list(sheet.iter_rows(values_only=True))
    if not rows:
        return []
    headers = tuple(str(value or "").strip() for value in rows[0][: len(MEME_LIBRARY_HEADERS)])
    if headers != MEME_LIBRARY_HEADERS:
        raise ValueError(f"表头必须严格为：{'、'.join(MEME_LIBRARY_HEADERS)}")
    return [
        {header: row[index] if index < len(row) else "" for index, header in enumerate(MEME_LIBRARY_HEADERS)}
        for row in rows[1:]
        if any(str(value or "").strip() for value in row[: len(MEME_LIBRARY_HEADERS)])
    ]


def _csv_rows(content: bytes) -> list[dict[str, Any]]:
    try:
        text = content.decode("utf-8-sig")
    except UnicodeDecodeError as exc:
        raise ValueError("CSV 必须使用 UTF-8 编码") from exc
    reader = csv.DictReader(io.StringIO(text))
    headers = tuple(str(value or "").strip() for value in (reader.fieldnames or []))
    if headers != MEME_LIBRARY_HEADERS:
        raise ValueError(f"表头必须严格为：{'、'.join(MEME_LIBRARY_HEADERS)}")
    return list(reader)


def parse_meme_library_file(filename: str, content: bytes) -> tuple[list[dict[str, Any]], list[str]]:
    if not content:
        raise ValueError("上传文件为空")
    if len(content) > MAX_MEME_IMPORT_BYTES:
        raise ValueError("热梗库文件不能超过 8 MB")
    suffix = Path(filename or "").suffix.lower()
    if suffix == ".xlsx":
        raw_rows = _xlsx_rows(content)
    elif suffix == ".csv":
        raw_rows = _csv_rows(content)
    else:
        raise ValueError("只支持 .xlsx 或 UTF-8 .csv 文件")
    if len(raw_rows) > MAX_MEME_IMPORT_ROWS:
        raise ValueError(f"单次最多导入 {MAX_MEME_IMPORT_ROWS} 条")

    entries: list[dict[str, Any]] = []
    errors: list[str] = []
    seen: set[str] = set()
    for row_number, row in enumerate(raw_rows, start=2):
        try:
            entry = _normalize_import_row(row, row_number)
            if not entry["normalized_phrase"] or entry["normalized_phrase"] in seen:
                raise ValueError(f"第 {row_number} 行热梗为空或在文件内重复")
            seen.add(entry["normalized_phrase"])
            entry["retrieval_text"] = build_meme_retrieval_text(entry)
            entry["content_hash"] = build_meme_content_hash(entry)
            entries.append(entry)
        except ValueError as exc:
            errors.append(str(exc))
    return entries, errors[:50]


def import_user_meme_entries(
    db: Session,
    *,
    owner: User,
    filename: str,
    content: bytes,
) -> dict[str, Any]:
    entries, errors = parse_meme_library_file(filename, content)
    namespace = f"user:{owner.id}"
    existing = {
        item.normalized_phrase: item
        for item in db.scalars(
            select(MemeEntry).where(MemeEntry.namespace == namespace)
        ).all()
    }
    imported = 0
    updated = 0
    skipped = 0
    for payload in entries:
        item = existing.get(payload["normalized_phrase"])
        if item is None:
            item = MemeEntry(
                owner_id=owner.id,
                namespace=namespace,
                source_type="user",
                review_status="approved",
                enabled=True,
                library_version="user-import",
                metadata_payload={"source_file": Path(filename).name},
                **payload,
            )
            db.add(item)
            existing[item.normalized_phrase] = item
            imported += 1
            continue
        if item.content_hash == payload["content_hash"]:
            skipped += 1
            continue
        for key, value in payload.items():
            setattr(item, key, value)
        item.enabled = True
        item.review_status = "approved"
        item.embedding = None
        item.embedding_model = ""
        item.metadata_payload = {"source_file": Path(filename).name}
        updated += 1
    db.commit()
    return {
        "imported": imported,
        "updated": updated,
        "skipped": skipped,
        "total": len(entries),
        "errors": errors,
    }


def visible_meme_filters(owner_id: UUID) -> tuple[Any, ...]:
    return (
        or_(MemeEntry.namespace == "builtin", MemeEntry.owner_id == owner_id),
        MemeEntry.enabled.is_(True),
        MemeEntry.review_status == "approved",
    )


def index_visible_meme_entries(
    db: Session,
    *,
    owner: User,
    force: bool = False,
) -> dict[str, Any]:
    config = build_embedding_config(owner.preferences)
    visible = db.scalars(
        select(MemeEntry)
        .where(*visible_meme_filters(owner.id))
        .order_by(MemeEntry.source_type, MemeEntry.phrase)
    ).all()
    if config is None:
        return {
            "status": "unavailable",
            "indexed": 0,
            "total_visible": len(visible),
            "embedding_model": "",
            "message": "请先在设置中开启并配置 Qwen Embedding",
        }
    for item in visible:
        retrieval_text = build_meme_retrieval_text(
            {
                "phrase": item.phrase,
                "meaning": item.meaning,
                "suitable_scenes": item.suitable_scenes,
            }
        )
        if item.retrieval_text != retrieval_text:
            item.retrieval_text = retrieval_text
            item.embedding = None
            item.embedding_model = ""
    pending = [
        item
        for item in visible
        if (
            force
            or item.embedding is None
            or item.embedding_model != meme_embedding_model_key(config.model)
        )
    ]
    if not pending:
        return {
            "status": "completed",
            "indexed": 0,
            "total_visible": len(visible),
            "embedding_model": config.model,
            "message": "热梗向量索引已是最新",
        }

    client = EmbeddingClient(config)
    indexed = 0
    batch_size = max(1, min(int(settings.embedding_batch_size), 20))
    for start in range(0, len(pending), batch_size):
        batch = pending[start : start + batch_size]
        vectors = client.embed(
            [item.retrieval_text for item in batch],
            text_type="document",
            instruct=(
                "Represent verified Chinese internet expressions for retrieval by "
                "speaker relationship, conversational intent, emotion, and scene."
            ),
        )
        for item, vector in zip(batch, vectors, strict=True):
            item.embedding = vector
            item.embedding_model = meme_embedding_model_key(config.model)
        db.commit()
        indexed += len(batch)
    return {
        "status": "completed",
        "indexed": indexed,
        "total_visible": len(visible),
        "embedding_model": config.model,
        "message": f"已更新 {indexed} 条热梗向量",
    }
